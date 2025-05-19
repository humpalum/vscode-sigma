import json
import csv
import requests
import openpyxl
import os

finobj = []
mitre_version = "v17.1"

# Get this at https://github.com/mitre/cti/blob/master/enterprise-attack/enterprise-attack.json
#with open("enterprise-attack.json") as f:
#    data = json.load(f)
#    for i in data["objects"]:
        # if i["type"] == "attack-pattern":
#        if "external_references" in i:
#            for id in i["external_references"]:
#                if "external_id" in id and "description" in i:
#                    finobj.append(
#                        {"name": i["name"], "description": i["description"], "tag": id["external_id"], "url": id["url"]})


def download_csv_xlsx_file(url):
    file_name = url.split("/")[-1]
    response = requests.get(url)
    if response.status_code == 200:
        with open(file_name, "wb") as file:
            file.write(response.content)
            return file_name
    else:
        print(f"Failed to download the file. Status code: {response.status_code}")

def convert_xlsx_to_csv(xlsx_file):
    workbook = openpyxl.load_workbook(xlsx_file)
    sheet_name = workbook.sheetnames[0]  # Use the first sheet
    sheet = workbook[sheet_name]
    csv_file = xlsx_file.replace(".xlsx", ".csv")
    with open(csv_file, "w", newline="") as file:
        writer = csv.writer(file)
        for row in sheet.iter_rows(values_only=True):
            writer.writerow(row)
    if csv_file:
        os.remove(xlsx_file)
        return csv_file
    else:
        print(f"Failed to convert {xlsx_file} to CSV.")
        return None

################## collecting tactics ##################
# Get this at https://attack.mitre.org/docs/enterprise-attack-v17.1/enterprise-attack-v17.1-tactics.xlsx
# Then save as csv the first sheet
tactic_file = download_csv_xlsx_file(f"https://attack.mitre.org/docs/enterprise-attack-{mitre_version}/enterprise-attack-{mitre_version}-tactics.xlsx")
tactic_csv_file = convert_xlsx_to_csv(tactic_file)
if tactic_csv_file:
    with open(tactic_csv_file, mode="r") as f:
        data = csv.reader(f, delimiter=',')
        for row in data:
            id = row[0]
            name = row[2]
            desc = row[3]
            url = row[4]
            if id == "ID":   # skipping header
                continue
            finobj.append({"name": name, "description": desc, "tag":id, "url":url})
else:
    Exception("Failed to convert tactics file to CSV.")


################## collecting techniques ##################
# Get this at https://attack.mitre.org/docs/enterprise-attack-v17.1/enterprise-attack-v17.1-techniques.xlsx
# Then save as csv the first sheet
technique_file = download_csv_xlsx_file(f"https://attack.mitre.org/docs/enterprise-attack-{mitre_version}/enterprise-attack-{mitre_version}-techniques.xlsx")
technique_csv_file = convert_xlsx_to_csv(technique_file)
if technique_csv_file:
    with open(technique_csv_file, mode="r") as f:
        data = csv.reader(f, delimiter=',')
        for row in data:
            id = row[0]
            name = row[2]
            desc = row[3]
            url = row[4]
            if id == "ID": # skipping header
                continue
            finobj.append({"name": name, "description": desc, "tag":id, "url":url})
else:
    Exception("Failed to convert techniques file to CSV.")

################## collecting software ##################
# Get this at https://attack.mitre.org/docs/enterprise-attack-v17.1/enterprise-attack-v17.1-software.xlsx
# Then save as csv the first sheet
software_file = download_csv_xlsx_file(f"https://attack.mitre.org/docs/enterprise-attack-{mitre_version}/enterprise-attack-{mitre_version}-software.xlsx")
software_csv_file = convert_xlsx_to_csv(software_file)
if software_csv_file:
    with open(software_csv_file, mode="r") as f:
        data = csv.reader(f, delimiter=',')
        for row in data:
            id = row[0]
            name = row[2]
            desc = row[3]
            url = row[4]
            if id == "ID": # skipping header
                continue
            finobj.append({"name": name, "description": desc, "tag":id, "url":url})
else:
    Exception("Failed to convert software file to CSV.")


################## collecting groups ##################
# Get this at https://attack.mitre.org/docs/enterprise-attack-v17.1/enterprise-attack-v17.1-groups.xlsx
# Then save as csv the first sheet
groups_file = download_csv_xlsx_file(f"https://attack.mitre.org/docs/enterprise-attack-{mitre_version}/enterprise-attack-{mitre_version}-groups.xlsx")
groups_csv_file = convert_xlsx_to_csv(groups_file)
if groups_csv_file:
    with open(groups_csv_file, mode="r") as f:
        data = csv.reader(f, delimiter=',')
        for row in data:
            id = row[0]
            name = row[2]
            desc = row[3]
            url = row[4]
            if id == "ID": # skipping header
                continue
            finobj.append({"name": name, "description": desc, "tag":id, "url":url})
else:
    Exception("Failed to convert groups file to CSV.")

            
################## collecting campaigns ##################
# Get this at https://attack.mitre.org/docs/enterprise-attack-v17.1/enterprise-attack-v17.1-campaigns.xlsx
# Then save as csv the first sheet
campaigns_file = download_csv_xlsx_file(f"https://attack.mitre.org/docs/enterprise-attack-{mitre_version}/enterprise-attack-{mitre_version}-campaigns.xlsx")
campaigns_csv_file = convert_xlsx_to_csv(campaigns_file)
if campaigns_csv_file:
    with open(campaigns_csv_file, mode="r") as f:
        data = csv.reader(f, delimiter=',')
        for row in data:
            id = row[0]
            name = row[2]
            desc = row[3]
            url = row[4]
            if id == "ID": # skipping header
                continue
            finobj.append({"name": name, "description": desc, "tag":id, "url":url})
            

################## collecting datasources ##################
# Get this at https://attack.mitre.org/docs/enterprise-attack-v17.1/enterprise-attack-v17.1-datasources.xlsx
# Then save as csv the first sheet
datasources_file = download_csv_xlsx_file(f"https://attack.mitre.org/docs/enterprise-attack-{mitre_version}/enterprise-attack-{mitre_version}-datasources.xlsx")
datasources_csv_file = convert_xlsx_to_csv(datasources_file)
if datasources_csv_file:
    with open(datasources_csv_file, mode="r") as f:
        data = csv.reader(f, delimiter=',')
        for row in data:
            id = row[1]
            name = row[0]
            desc = row[3]
            url = row[10]
            if id == "" or id == "ID": # skipping header and empty datasource
                continue
            finobj.append({"name": name, "description": desc, "tag":id, "url":url})
else:
    Exception("Failed to convert datasources file to CSV.")
            
################## collecting mitigations ##################
# Get this at https://attack.mitre.org/docs/enterprise-attack-v17.1/enterprise-attack-v17.1-mitigations.xlsx
# Then save as csv the first sheet
mitigations_file = download_csv_xlsx_file(f"https://attack.mitre.org/docs/enterprise-attack-{mitre_version}/enterprise-attack-{mitre_version}-mitigations.xlsx")
mitigations_csv_file = convert_xlsx_to_csv(mitigations_file)
if mitigations_csv_file:
    with open(mitigations_csv_file, mode="r") as f:
        data = csv.reader(f, delimiter=',')
        for row in data:
            id = row[0]
            name = row[2]
            desc = row[3]
            url = row[4]
            if id == "ID": # skipping header
                continue
            finobj.append({"name": name, "description": desc, "tag":id, "url":url})
else:
    Exception("Failed to convert mitigations file to CSV.")

# Get this at https://d3fend.mitre.org/ontologies/d3fend.csv
defend_file = download_csv_xlsx_file("https://d3fend.mitre.org/ontologies/d3fend.csv")
with open(defend_file) as file:
    spamreader = csv.reader(file, delimiter=',')
    for row in spamreader:
        id = row[0]
        name= row[2] if row[2] != '' else row[3] if row[3] != '' else row[4]
        desc = row[5]
        finobj.append({"name": name, "description": desc, "tag":id, "url":"https://d3fend.mitre.org/"})


finobj.sort(key=lambda x: x["tag"])
with open("./src/techniques.json", 'w') as nf:
    nf.write(json.dumps(finobj, indent=4))
